"""Turn an uploaded CSV or Excel (.xlsx) file into plain rows of text. Nothing is stored.

This is the one place that touches an untrusted file, so it is deliberately strict:
* only .csv and .xlsx (no .xls, no macro workbooks), at most 1 MB, at most 300 data rows;
* Excel files are opened read-only with data_only=True, so a formula is never evaluated: the
  cell's last saved value is read (or nothing, which the validator reports as a missing value);
* only the first sheet is read;
* a zip-bomb guard refuses archives that would unpack to far more than the upload's size;
* the XML is parsed by defusedxml (openpyxl uses it automatically when it is installed).

It knows nothing about societies or flats: it returns text, and the validator decides if the
text makes sense.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

MAX_FILE_BYTES = 1_000_000
MAX_ROWS = 300
MAX_UNZIPPED_BYTES = 10_000_000
MAX_ZIP_ENTRIES = 200

# What we call each column, and the headings people actually type for it (compared after
# lower-casing and replacing punctuation with spaces, so "Flat No." and "flat_no" both match).
COLUMN_ALIASES = {
    "flat": {"flat", "flat no", "flat number", "flatno", "flat num", "unit", "house", "house no"},
    "name": {"name", "tenant", "tenant name", "resident", "resident name", "owner", "full name"},
    "phone": {
        "phone",
        "phone no",
        "phone number",
        "mobile",
        "mobile no",
        "mobile number",
        "contact",
    },
    "email": {"email", "e mail", "email id", "email address"},
    "maintenance": {
        "maintenance",
        "maintenance amount",
        "maintenance fee",
        "maintenance fees",
        "amount",
        "fee",
        "fees",
    },
    "role": {"role"},
}
REQUIRED = ("flat", "name", "maintenance")
HEADING_HELP = {
    "flat": "flat number",
    "name": "tenant name",
    "maintenance": "maintenance amount",
}


class ImportFileError(Exception):
    """The file itself cannot be used; the message says why in plain words."""


@dataclass
class ParsedFile:
    rows: list  # [(spreadsheet row number, {column: text})]
    ignored_columns: list  # headings we did not recognise (reported, never used)


def _heading(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def _text(value) -> str:
    """One cell as clean text. Numbers lose a trailing '.0' (Excel stores 9876543210 as a float)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _csv_rows(data: bytes) -> list:
    for encoding in ("utf-8-sig", "cp1252"):  # Excel's plain "CSV" is often cp1252, not UTF-8
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError(
            "The file's text could not be read. Save it as 'CSV UTF-8' and retry."
        )
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    try:
        return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    except csv.Error as exc:
        raise ImportFileError(f"The file could not be read as CSV ({exc}).") from None


def _xlsx_rows(data: bytes) -> list:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ImportFileError("This does not look like a real Excel (.xlsx) file.") from None
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNZIPPED_BYTES:
            raise ImportFileError("This Excel file is too large or unusual to read safely.")
        if any(e.filename.lower().startswith("xl/vbaproject") for e in entries):
            raise ImportFileError("Excel files with macros are not accepted. Save a plain .xlsx.")
    from openpyxl import load_workbook  # imported here so CSV-only paths never load it

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 50)]
        finally:
            workbook.close()
    except ImportFileError:
        raise
    except Exception:  # noqa: BLE001 - a third-party parser boundary: any failure is "unreadable"
        raise ImportFileError(
            "The Excel file could not be read. Is it damaged or password protected?"
        ) from None


def read_rows(filename: str, data: bytes) -> ParsedFile:
    """Parse an upload. Raises ImportFileError with a plain-language reason when unusable."""
    name = (filename or "").lower()
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("The file is larger than 1 MB. A flats list should be far smaller.")
    if name.endswith(".csv"):
        table = _csv_rows(data)
    elif name.endswith(".xlsx"):
        table = _xlsx_rows(data)
    else:
        raise ImportFileError(
            "Upload a .csv or .xlsx file. Older .xls and macro files are not accepted."
        )

    # The first non-empty row is the heading row.
    header_at = next((i for i, row in enumerate(table) if any(_text(c) for c in row)), None)
    if header_at is None:
        raise ImportFileError("The file is empty.")

    wanted = {}  # column index -> our name for it
    ignored = []
    for index, cell in enumerate(table[header_at]):
        heading = _heading(_text(cell))
        if not heading:
            continue
        ours = next((key for key, names in COLUMN_ALIASES.items() if heading in names), None)
        if ours is None or ours in wanted.values():
            ignored.append(_text(cell))
        else:
            wanted[index] = ours

    missing = [HEADING_HELP[key] for key in REQUIRED if key not in wanted.values()]
    if missing:
        raise ImportFileError(
            "The first row must name these columns: " + ", ".join(missing) + ". "
            "Download the template to see the expected headings."
        )

    rows = []
    for offset, raw in enumerate(table[header_at + 1 :], start=header_at + 2):
        values = {ours: _text(raw[i]) if i < len(raw) else "" for i, ours in wanted.items()}
        if any(values.values()):  # blank lines between rows are skipped
            rows.append((offset, values))
    if not rows:
        raise ImportFileError("The file has headings but no rows.")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(
            f"The file has more than {MAX_ROWS} rows. Split it into smaller files."
        )
    return ParsedFile(rows=rows, ignored_columns=ignored)
